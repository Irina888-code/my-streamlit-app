import streamlit as st
import requests
import os
import pandas as pd
from datetime import datetime
try:
    from docx import Document
except ImportError:
    import subprocess
    subprocess.check_call(['pip', 'install', 'python-docx'])
    from docx import Document

st.set_page_config(
    page_title="ИИ-помощник для семейного канала",
    page_icon="👨‍👩‍👧‍👦",
    layout="wide",
    initial_sidebar_state="expanded"
)

def get_api_key(key_name):
    """Безопасное получение API ключей с улучшенной отладкой"""
    # Пробуем получить из secrets
    try:
        if key_name in st.secrets:
            value = st.secrets[key_name]
            if value and len(str(value)) > 5:
                return str(value).strip()
            else:
                st.warning(f"Ключ {key_name} в secrets пустой или слишком короткий")
                return ""
    except Exception as e:
        st.warning(f"Ошибка доступа к secrets для {key_name}: {e}")
    
    # Пробуем из переменных окружения
    env_value = os.getenv(key_name, "")
    if env_value:
        return env_value.strip()
    
    return ""

# Отладочная информация (можно удалить после проверки)
st.sidebar.markdown("---")
st.sidebar.markdown("### 🔍 Отладка")
try:
    available_keys = list(st.secrets.keys()) if hasattr(st, 'secrets') else []
    st.sidebar.write(f"Доступные ключи: {available_keys}")
except:
    st.sidebar.write("Secrets не доступны")

CSS_STYLES = """
<style>
#MainMenu {visibility: hidden;}
.stAppDeployButton {display: none;}

.main-title {
    background: linear-gradient(135deg, #667eea 0%, #764ba2 50%, #f093fb 100%);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
    background-clip: text;
    font-size: 3.2em;
    font-weight: bold;
    text-align: center;
    margin-bottom: 10px;
}   

.subtitle {
    text-align: center;
    color: #764ba2;
    font-size: 1.2em;
    margin-bottom: 30px;
    font-weight: 500;
}

.stButton > button {
    border-radius: 10px;
    font-weight: 500;
    width: 100%;
}

[data-testid="stSidebar"] {
    background: linear-gradient(180deg, #667eea 0%, #764ba2 100%);
}
[data-testid="stSidebar"] * {
    color: white !important;
}
[data-testid="stSidebar"] input,
[data-testid="stSidebar"] select {
    color: white !important;
}
[data-testid="stSidebar"] .stSelectbox > div > div > select {
    background-color: rgba(255,255,255,0.1) !important;
    color: white !important;
}
.sidebar-header {
    background: rgba(255,255,255,0.15);
    color: white;
    padding: 15px;
    border-radius: 10px;
    text-align: center;
    margin-bottom: 20px;
}
</style>
"""

st.markdown(CSS_STYLES, unsafe_allow_html=True)

def build_prompt(topic, platform, tone, temperature, text_length):
    prompt_text = (
        f"Persona: Опытный контент-мейкер и копирайтер для семейных медиа. Эксперт по созданию лаконичных, практичных и вовлекающих материалов для родителей детей 2-10 лет.\n\n"
        f"Task: Написать готовый к публикации пост для платформы {platform} на тему: '{topic}'.\n\n"
        f"Context: Аудитория - родители дошкольников и младших школьников (2-10 лет).\n\n"
        f"Format: Структура: 1) Вступление - 2) 2-3 практических совета - 3) Вопрос к читателям. Эмодзи умеренно. В конце 8-10 хештегов включая #семья #родители #воспитание #дети #семейноевремя.\n\n"
        f"Критерии:\n"
        f"- Без воды и клише\n"
        f"- Короткие абзацы (1-3 строки)\n"
        f"- Креативность: {temperature}\n"
        f"- Объём: {text_length}\n"
        f"- Практичные советы для возраста 2-10 лет\n\n"
        f"Выводить только финальный пост."
    )
    return prompt_text

def call_yandex(prompt):
    api_key = get_api_key("YANDEX_API_KEY")
    folder_id = get_api_key("YANDEX_FOLDER_ID")
    
    st.sidebar.write(f"Yandex API Key получен: {'Да' if api_key else 'Нет'}, длина: {len(api_key) if api_key else 0}")
    st.sidebar.write(f"Yandex Folder ID получен: {'Да' if folder_id else 'Нет'}, длина: {len(folder_id) if folder_id else 0}")
    
    if not api_key or not folder_id:
        return "⚠️ Не настроен YandexGPT. Добавьте YANDEX_API_KEY и YANDEX_FOLDER_ID в Secrets\n\n📝 Инструкция:\n1. В боковой панели нажмите 'Manage app'\n2. Перейдите в раздел 'Secrets'\n3. Добавьте:\n   YANDEX_API_KEY = \"ваш_ключ\"\n   YANDEX_FOLDER_ID = \"ваш_folder_id\""
    
    url = "https://llm.api.cloud.yandex.net/foundationModels/v1/completion"
    headers = {
        "Authorization": f"Api-Key {api_key}",
        "x-folder-id": folder_id,
        "Content-Type": "application/json"
    }
    payload = {
        "modelUri": f"gpt://{folder_id}/yandexgpt-lite",
        "completionOptions": {
            "stream": False, 
            "temperature": 0.7, 
            "maxTokens": 2000
        },
        "messages": [
            {"role": "system", "text": "Ты опытный контент-мейкер для семейного блога. Отвечай только на русском языке."},
            {"role": "user", "text": prompt}
        ]
    }
    
    try:
        with st.spinner("🔄 Запрос к YandexGPT..."):
            resp = requests.post(url, headers=headers, json=payload, timeout=60)
            
            if resp.status_code == 401:
                return "❌ Ошибка авторизации YandexGPT. Проверьте правильность API ключа и Folder ID"
            elif resp.status_code == 403:
                return "❌ Доступ запрещен. Проверьте аккаунт Yandex Cloud"
            elif resp.status_code == 429:
                return "❌ Превышен лимит запросов к YandexGPT"
            
            resp.raise_for_status()
            result = resp.json()
            
            # Проверяем структуру ответа
            if "result" in result and "alternatives" in result["result"] and len(result["result"]["alternatives"]) > 0:
                return result["result"]["alternatives"][0]["message"]["text"]
            else:
                return f"❌ Неожиданный формат ответа YandexGPT: {result}"
                
    except requests.exceptions.Timeout:
        return "❌ Таймаут YandexGPT (60 сек). Попробуйте ещё раз."
    except requests.exceptions.RequestException as e:
        return f"❌ Ошибка сети YandexGPT: {str(e)}"
    except Exception as e:
        return f"❌ Ошибка YandexGPT: {str(e)}"

def call_deepseek(prompt):
    api_key = get_api_key("DEEPSEEK_API_KEY")
    
    st.sidebar.write(f"DeepSeek API Key получен: {'Да' if api_key else 'Нет'}, длина: {len(api_key) if api_key else 0}")
    
    if not api_key:
        return "⚠️ Не указан DEEPSEEK_API_KEY. Добавьте ключ в Secrets\n\n📝 Инструкция:\n1. В боковой панели нажмите 'Manage app'\n2. Перейдите в раздел 'Secrets'\n3. Добавьте: DEEPSEEK_API_KEY = \"sk-ваш_ключ\""
    
    url = "https://api.deepseek.com/v1/chat/completions"
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json"
    }
    payload = {
        "model": "deepseek-chat",
        "messages": [
            {"role": "system", "content": "Ты опытный контент-мейкер для семейного блога. Отвечай только на русском языке."},
            {"role": "user", "content": prompt}
        ],
        "max_tokens": 2000,
        "temperature": 0.7
    }
    
    try:
        with st.spinner("🔄 Запрос к DeepSeek..."):
            resp = requests.post(url, headers=headers, json=payload, timeout=60)
            
            if resp.status_code == 401:
                return "❌ Ошибка авторизации DeepSeek. Проверьте правильность API ключа"
            elif resp.status_code == 429:
                return "❌ Превышен лимит запросов к DeepSeek"
            
            resp.raise_for_status()
            result = resp.json()
            
            if "choices" in result and len(result["choices"]) > 0:
                return result["choices"][0]["message"]["content"]
            else:
                return f"❌ Неожиданный формат ответа DeepSeek: {result}"
                
    except requests.exceptions.Timeout:
        return "❌ Таймаут DeepSeek (60 сек). Попробуйте ещё раз."
    except requests.exceptions.RequestException as e:
        return f"❌ Ошибка сети DeepSeek: {str(e)}"
    except Exception as e:
        return f"❌ Ошибка DeepSeek: {str(e)}"

def save_to_word(text, filename):
    doc = Document()
    doc.add_heading("Пост для семейного канала", 0)
    doc.add_paragraph(text)
    doc.save(filename)
    return filename

def save_to_history(topic, platform, tone, text_length, yandex_text, deepseek_text):
    try:
        data = {
            "Дата": [datetime.now().strftime("%d.%m.%Y %H:%M")],
            "Тема": [topic],
            "Платформа": [platform],
            "Тон": [tone],
            "Длина": [text_length],
            "YandexGPT": [yandex_text[:100] + "..." if len(yandex_text) > 100 else yandex_text],
            "DeepSeek": [deepseek_text[:100] + "..." if len(deepseek_text) > 100 else deepseek_text]
        }
        df = pd.DataFrame(data)
        csv_file = "history.csv"
        
        if os.path.exists(csv_file):
            df.to_csv(csv_file, mode="a", header=False, index=False, encoding='utf-8-sig')
        else:
            df.to_csv(csv_file, index=False, encoding='utf-8-sig')
        return True
    except Exception as e:
        st.error(f"Ошибка сохранения: {e}")
        return False

# Инициализация session_state
if "yandex_text" not in st.session_state:
    st.session_state.yandex_text = ""
if "deepseek_text" not in st.session_state:
    st.session_state.deepseek_text = ""
if "topic" not in st.session_state:
    st.session_state.topic = ""
if "platform" not in st.session_state:
    st.session_state.platform = "Telegram"
if "tone" not in st.session_state:
    st.session_state.tone = "Тёплый и искренний"
if "text_length_display" not in st.session_state:
    st.session_state.text_length_display = "300-400 слов"
if "generated" not in st.session_state:
    st.session_state.generated = False

# Основной интерфейс
st.markdown("<div class='main-title'>👨‍👩‍👧‍👦 ИИ-помощник для семейного канала</div>", unsafe_allow_html=True)
st.markdown("<div class='subtitle'>✨ Генерация постов с хештегами и эмодзи через YandexGPT и DeepSeek</div>", unsafe_allow_html=True)

quick_topics = [
    "5 идей для семейных выходных",
    "Как пережить утренние сборы в детский сад без слёз",
    "Топ-10 настольных игр для детей 3-7 лет",
    "Рецепты полезных перекусов для школы",
    "Семейные традиции, которые сближают",
    "Как научить ребёнка читать через игру",
    "Идеи для домашнего праздника ребёнка",
    "Как справиться с детскими истериками в магазине"
]

col_a, col_b = st.columns([3, 1])
with col_a:
    topic_input = st.text_input("Тема поста", placeholder="Введите свою тему", value=st.session_state.topic)
with col_b:
    selected_quick = st.selectbox("Или выберите:", ["- Готовые темы -"] + quick_topics)
    if selected_quick != "- Готовые темы -":
        topic_input = selected_quick

col1, col2, col3 = st.columns(3)
with col1:
    platform_input = st.selectbox("Платформа", ["Telegram", "ВКонтакте", "Дзен", "Instagram"])
with col2:
    tone_input = st.selectbox("Тон", ["Тёплый и искренний", "Экспертный", "Мотивирующий", "Лёгкий и юмористический"])
with col3:
    temperature = st.slider("Креативность", 0.1, 1.0, 0.7, 0.1)

# Получаем настройки из боковой панели
model_option = st.session_state.get("model_option", "YandexGPT + DeepSeek")
text_length_display = st.session_state.get("text_length_display", "300-400 слов")

# Кнопка генерации
if st.button("🚀 Сгенерировать посты", type="primary", use_container_width=True):
    if not topic_input:
        st.error("❌ Введите тему поста!")
    else:
        with st.spinner("🤖 ИИ генерирует посты... Подождите 15-30 секунд"):
            # Преобразуем длину текста
            length_map = {
                "Менее 300 слов": "менее 300 слов",
                "300-400 слов": "300-400 слов", 
                "500-600 слов": "500-600 слов",
                "Более 600 слов": "более 600 слов"
            }
            text_length_value = length_map.get(text_length_display, "300-400 слов")
            
            prompt = build_prompt(topic_input, platform_input, tone_input, temperature, text_length_value)
            
            # Генерируем в зависимости от выбранной модели
            yandex_result = ""
            deepseek_result = ""
            
            if model_option in ["YandexGPT + DeepSeek", "Только YandexGPT"]:
                yandex_result = call_yandex(prompt)
            else:
                yandex_result = "❌ Выбран режим 'Только DeepSeek'"
                
            if model_option in ["YandexGPT + DeepSeek", "Только DeepSeek"]:
                deepseek_result = call_deepseek(prompt)
            else:
                deepseek_result = "❌ Выбран режим 'Только YandexGPT'"
            
            # Сохраняем в session_state
            st.session_state.yandex_text = yandex_result
            st.session_state.deepseek_text = deepseek_result
            st.session_state.topic = topic_input
            st.session_state.platform = platform_input
            st.session_state.tone = tone_input
            st.session_state.generated = True
            
            # Сохраняем в историю только если есть успешные ответы
            if ("Не настроен" not in yandex_result and "Не указан" not in deepseek_result and 
                not yandex_result.startswith("❌") and not deepseek_result.startswith("❌")):
                save_to_history(topic_input, platform_input, tone_input, text_length_display, 
                              yandex_result, deepseek_result)
            
            st.success("✅ Посты успешно сгенерированы!")
            st.rerun()

# Отображаем результаты, если есть
if st.session_state.generated and (st.session_state.yandex_text or st.session_state.deepseek_text):
    st.markdown("---")
    
    # Создаем вкладки
    tab1, tab2, tab3, tab4 = st.tabs(["📝 YandexGPT", "🤖 DeepSeek", "📱 Предпросмотр", "📊 Сравнение"])
    
    with tab1:
        if st.session_state.yandex_text:
            st.subheader("YandexGPT (Lite)")
            
            # Проверяем на ошибки
            if "Не настроен" in st.session_state.yandex_text or "Не указан" in st.session_state.yandex_text:
                st.error(st.session_state.yandex_text)
            else:
                st.markdown(st.session_state.yandex_text)
                
                col_act1, col_act2 = st.columns(2)
                with col_act1:
                    try:
                        doc_file = save_to_word(st.session_state.yandex_text, "yandex_post.docx")
                        with open(doc_file, "rb") as f:
                            st.download_button("📥 Скачать Word", f, file_name="post_yandex.docx", 
                                             mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
                    except:
                        st.error("Ошибка создания Word файла")
                with col_act2:
                    st.code(st.session_state.yandex_text, language="markdown", line_numbers=False)
        else:
            st.info("Нет сгенерированного текста от YandexGPT")
    
    with tab2:
        if st.session_state.deepseek_text:
            st.subheader("DeepSeek")
            
            # Проверяем на ошибки
            if "Не настроен" in st.session_state.deepseek_text or "Не указан" in st.session_state.deepseek_text:
                st.error(st.session_state.deepseek_text)
            else:
                st.markdown(st.session_state.deepseek_text)
                
                col_act1, col_act2 = st.columns(2)
                with col_act1:
                    try:
                        doc_file = save_to_word(st.session_state.deepseek_text, "deepseek_post.docx")
                        with open(doc_file, "rb") as f:
                            st.download_button("📥 Скачать Word", f, file_name="post_deepseek.docx",
                                             mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
                    except:
                        st.error("Ошибка создания Word файла")
                with col_act2:
                    st.code(st.session_state.deepseek_text, language="markdown", line_numbers=False)
        else:
            st.info("Нет сгенерированного текста от DeepSeek")
    
    with tab3:
        st.subheader("Предпросмотр в стиле Telegram")
        
        if st.session_state.yandex_text and "Не настроен" not in st.session_state.yandex_text and "Не указан" not in st.session_state.yandex_text:
            st.markdown("**📨 От YandexGPT:**")
            preview_ya = st.session_state.yandex_text[:400] + "..." if len(st.session_state.yandex_text) > 400 else st.session_state.yandex_text
            st.markdown(f"""
            <div style='background: #f0f2f6; border-radius: 15px; padding: 15px; max-width: 500px; margin: 10px 0;'>
                <div style='background: white; border-radius: 12px; padding: 15px; font-size: 14px; line-height: 1.5;'>
                    {preview_ya.replace(chr(10), '<br>')}
                </div>
                <div style='text-align: right; color: #666; font-size: 11px; margin-top: 8px;'>
                    {datetime.now().strftime("%H:%M")}
                </div>
            </div>
            """, unsafe_allow_html=True)
        
        if st.session_state.deepseek_text and "Не настроен" not in st.session_state.deepseek_text and "Не указан" not in st.session_state.deepseek_text:
            st.markdown("**📨 От DeepSeek:**")
            preview_ds = st.session_state.deepseek_text[:400] + "..." if len(st.session_state.deepseek_text) > 400 else st.session_state.deepseek_text
            st.markdown(f"""
            <div style='background: #f0f2f6; border-radius: 15px; padding: 15px; max-width: 500px; margin: 10px 0;'>
                <div style='background: white; border-radius: 12px; padding: 15px; font-size: 14px; line-height: 1.5;'>
                    {preview_ds.replace(chr(10), '<br>')}
                </div>
                <div style='text-align: right; color: #666; font-size: 11px; margin-top: 8px;'>
                    {datetime.now().strftime("%H:%M")}
                </div>
            </div>
            """, unsafe_allow_html=True)
    
    with tab4:
        st.subheader("📊 Сравнение моделей")
        
        col_c1, col_c2 = st.columns(2)
        with col_c1:
            ya_len = len(st.session_state.yandex_text) if st.session_state.yandex_text else 0
            st.metric("📏 YandexGPT символов", ya_len)
        with col_c2:
            ds_len = len(st.session_state.deepseek_text) if st.session_state.deepseek_text else 0
            st.metric("📏 DeepSeek символов", ds_len)
        
        st.markdown("---")
        st.markdown(f"""
        **📌 Детали генерации:**
        - **Тема:** {st.session_state.topic}
        - **Платформа:** {st.session_state.platform}
        - **Тон:** {st.session_state.tone}
        - **Длина:** {text_length_display}
        
        **💡 Рекомендации:**
        - **YandexGPT** лучше для: коротких постов, простого русского языка, быстрой генерации
        - **DeepSeek** лучше для: развёрнутых текстов, креативных идей, сложных запросов
        """)
        
        # Кнопка очистки
        if st.button("🗑️ Очистить результаты"):
            st.session_state.generated = False
            st.session_state.yandex_text = ""
            st.session_state.deepseek_text = ""
            st.rerun()

# Боковая панель
with st.sidebar:
    st.markdown("<div class='sidebar-header'><h2>📋 Меню</h2></div>", unsafe_allow_html=True)
    
    st.write("### ⚙️ Настройки")
    
    model_option = st.selectbox(
        "🤖 Модель ИИ",
        ["YandexGPT + DeepSeek", "Только YandexGPT", "Только DeepSeek"],
        key="model_option"
    )
    
    text_length_display = st.selectbox(
        "📏 Длина текста",
        ["Менее 300 слов", "300-400 слов", "500-600 слов", "Более 600 слов"],
        key="text_length_display"
    )
    
    st.divider()
    
    st.write("### 📊 Статистика")
    if os.path.exists("history.csv"):
        try:
            df_history = pd.read_csv("history.csv", encoding='utf-8-sig')
            st.metric("📝 Всего постов", len(df_history))
            
            today = datetime.now().strftime("%d.%m.%Y")
            today_posts = df_history[df_history["Дата"].str.startswith(today)] if not df_history.empty else pd.DataFrame()
            st.metric("📅 Сегодня", len(today_posts))
        except Exception as e:
            st.info(f"История: {str(e)[:50]}")
    else:
        st.info("📭 История пуста")
    
    st.divider()
    
    st.write("### 📜 Последние посты")
    if os.path.exists("history.csv"):
        try:
            df_history = pd.read_csv("history.csv", encoding='utf-8-sig')
            if not df_history.empty:
                for idx, row in df_history.tail(3).iterrows():
                    with st.expander(f"📌 {row['Дата']}"):
                        st.write(f"**Тема:** {row['Тема'][:50]}{'...' if len(row['Тема']) > 50 else ''}")
                        st.write(f"**Платформа:** {row['Платформа']}")
                        st.write(f"**Тон:** {row['Тон']}")
            else:
                st.write("Нет сохранённых постов")
        except Exception as e:
            st.write("Ошибка загрузки истории")
    else:
        st.write("💡 Сохраните первый пост")
    
    st.divider()
    
    st.write("### 🔐 Статус API")
    
    # Проверяем наличие ключей
    yandex_key = get_api_key("YANDEX_API_KEY")
    yandex_folder = get_api_key("YANDEX_FOLDER_ID")
    deepseek_key = get_api_key("DEEPSEEK_API_KEY")
    
    if yandex_key and yandex_folder and len(yandex_key) > 10 and len(yandex_folder) > 5:
        st.success("✅ YandexGPT готов")
    else:
        st.error(f"❌ YandexGPT: {'нет ключей' if not yandex_key else f'Folder ID: {len(yandex_folder) if yandex_folder else 0} символов'}")
    
    if deepseek_key and len(deepseek_key) > 20:
        st.success("✅ DeepSeek готов")
    else:
        st.error(f"❌ DeepSeek: ключ {'отсутствует' if not deepseek_key else f'короткий ({len(deepseek_key)} символов)'}")
    
    st.divider()
    
    st.write("### ℹ️ О проекте")
    st.info("""
    **ИИ-помощник для семейного канала**
    
    Версия 3.1
    
    🔑 API ключи: настройте в Secrets
    
    📝 Генерирует посты для:
    • Telegram
    • ВКонтакте
    • Дзен
    • Instagram
    """)